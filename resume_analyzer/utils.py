import os
import re
import fitz
import docx

SKILL_CATEGORIES = {
    "web_development": ["html", "css", "javascript", "bootstrap", "react", "django"],
    "python_development": ["python", "django", "flask", "api", "backend"],
    "machine_learning": ["python", "machine learning", "deep learning", "tensorflow", "pytorch", "nlp"],
    "data_analyst": ["python", "sql", "data analysis", "pandas", "excel", "power bi"]
}

def extract_text_from_pdf(path):
    text = ""
    try:
        doc = fitz.open(path)
        for page in doc:
            text += page.get_text()
    except Exception as e:
        print(f"PDF extraction error: {e}")
        return ""
    return text

def extract_text_from_docx(path):
    try:
        doc = docx.Document(path)
        lines = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in lines:
                        lines.append(cell_text)
        return "\n".join(lines)
    except Exception as e:
        print(f"DOCX extraction error: {e}")
        return ""

def extract_text(path):
    if not path:
        return ""
    path_lower = path.lower()
    if path_lower.endswith(".pdf"):
        return extract_text_from_pdf(path)
    elif path_lower.endswith(".docx"):
        return extract_text_from_docx(path)
    return ""

def extract_skills(text):
    if not text:
        return []
    text = text.lower()
    skills = []
    all_skills = []
    for v in SKILL_CATEGORIES.values():
        all_skills.extend(v)

    for skill in set(all_skills):
        escaped = re.escape(skill)
        if re.search(r'(?<![a-zA-Z0-9])' + escaped + r'(?![a-zA-Z0-9])', text):
            skills.append(skill)

    return sorted(list(set(skills)))

def detect_domain(skills):
    if not skills:
        return "general"
    scores = {}
    skills_set = set(s.lower() for s in skills)
    for domain, keywords in SKILL_CATEGORIES.items():
        scores[domain] = len(skills_set & set(keywords))

    if all(v == 0 for v in scores.values()):
        return "general"

    return max(scores, key=scores.get)

def calculate_ats_score(skills, domain):
    if not domain or domain == "general" or domain not in SKILL_CATEGORIES:
        return 0

    domain_skills = SKILL_CATEGORIES[domain]
    if not domain_skills:
        return 0
    skills_set = set(s.lower() for s in skills)
    return int((len(skills_set & set(domain_skills)) / len(domain_skills)) * 100)

def get_missing_skills(found, domain):
    if not domain or domain == "general" or domain not in SKILL_CATEGORIES:
        return []

    found_set = set(s.lower() for s in found)
    return list(set(SKILL_CATEGORIES[domain]) - found_set)

def generate_suggestions(skills, missing, domain, ats):
    suggestions = []

    if ats >= 80:
        suggestions.append("Your resume is strong and well optimized for ATS systems.")
    elif ats >= 50:
        suggestions.append("Your resume is good but can be improved.")
    else:
        suggestions.append("Your resume needs improvement.")

    for m in missing:
        suggestions.append(f"Consider adding skill: {m}")

    if domain == "web_development":
        suggestions.append("Add React/Django projects.")
    elif domain == "machine_learning":
        suggestions.append("Add ML projects with real datasets.")
    elif domain == "data_analyst":
        suggestions.append("Include SQL dashboards or Power BI projects.")
    elif domain == "python_development":
        suggestions.append("Build API/backend projects.")

    return suggestions
